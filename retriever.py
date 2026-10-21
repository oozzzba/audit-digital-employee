"""检索层：两级知识库 + BM25 检索 + 元数据 + 可配置阈值

知识库分两级（与知识库规格一致）：
  - 公共知识库：审计法规库 / 风控规则库 / 行业案例库 / 提示词模板库
  - 阶段 Agent 专属知识库：法规库 / 业务规则库 / 模板库 / 案例库 / 工具库 / 术语库

每条知识片段携带元数据：来源 / 条款位置 / 生效期 / 置信度 / 所属库 / 阶段。
检索要求：返回「来源 + 条款位置 + 生效期 + 置信度」；过期（失效）不召回；引用须有依据；
Top K 默认 6、相似度阈值 0.75，均可在 RETRIEVAL_CONFIG 配置。

说明：MVIP 用「jieba 分词 + BM25 关键词检索」验证「检索→生成」主链路，
后续升级为「向量检索(embedding) + 混合检索」。
"""
import os
import re
import zipfile
import jieba
from rank_bm25 import BM25Okapi

KNOWLEDGE_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), "knowledge")
TEMPLATE_DIR = os.getenv("AUDIT_TEMPLATE_DIR", "/Users/mac/作品集/审计各阶段模版")
EXTERNAL_KB_DIR = os.getenv("AUDIT_EXTERNAL_KB_DIR", "/Users/mac/作品集/综合监督知识库")

# 检索配置（默认 Top K=6，相似度阈值 0.75，可配置）
RETRIEVAL_CONFIG = {
    "top_k": 6,
    "threshold": 0.75,  # 归一化后相似度阈值，低于此不召回
}

# 文档前缀 → 所属审计阶段。用于「阶段化检索范围」：
# 现场实施只检索法规/定性标准，报告编制检索报告规范，整改督办检索整改要点。
_STAGE_BY_PREFIX = {
    "01-": "field",
    "02-": "field",
    "03-": "field",
    "04-": "field",
    "05-": "report",
    "06-": "rectification",
    "07-": "preaudit",
    "08-": "field",
    "09-": "field",
    "10-": "initiation",
    "11-": "initiation",
    "12-": "initiation",
    "13-": "initiation",
    "14-": "initiation",
    "15-": "initiation",
    "16-": "initiation",
    "17-": "initiation",
    "18-": "initiation",
}

# 知识条目元数据（来源/条款位置/生效期/置信度/所属库）。
# 诚信约定：整篇知识用「全文」标注条款位置，不虚构具体条文号；生效期均为「现行有效」。
_DOC_META = {
    "01-审计法关键条款.md": {"article": "全文", "effective": "现行有效", "confidence": 0.95, "library": "公共·审计法规库"},
    "02-内部审计工作规定要点.md": {"article": "全文", "effective": "现行有效", "confidence": 0.95, "library": "公共·审计法规库"},
    "03-审计工作底稿模板.md": {"article": "全文", "effective": "现行有效", "confidence": 0.90, "library": "阶段专属·模板库"},
    "04-常见审计问题定性标准.md": {"article": "全文", "effective": "现行有效", "confidence": 0.93, "library": "公共·风控规则库"},
    "05-审计报告编制规范.md": {"article": "全文", "effective": "现行有效", "confidence": 0.90, "library": "阶段专属·模板库"},
    "06-审计整改督办要点.md": {"article": "全文", "effective": "现行有效", "confidence": 0.90, "library": "阶段专属·工具库"},
    "07-审计实施方案规范.md": {"article": "全文", "effective": "现行有效", "confidence": 0.90, "library": "阶段专属·模板库"},
    "08-内部审计制度依据.md": {"article": "全文", "effective": "现行有效", "confidence": 0.95, "library": "公共·审计法规库"},
    "09-工程财务典型问题与案例.md": {"article": "全文", "effective": "现行有效", "confidence": 0.85, "library": "公共·行业案例库"},
    "10-审计项目类型库.md": {"article": "类型枚举与判定规则", "effective": "现行有效", "confidence": 0.96, "library": "立项专属·审计项目类型库"},
    "11-历史立项指令样本库.md": {"article": "口语化指令样本", "effective": "现行有效", "confidence": 0.88, "library": "立项专属·历史立项指令样本库"},
    "12-多源业务数据元数据索引库.md": {"article": "六类业务数据字段语义", "effective": "现行有效", "confidence": 0.93, "library": "立项专属·多源业务数据元数据索引库"},
    "13-历史审计台账风险库.md": {"article": "历史风险模式与处置", "effective": "现行有效", "confidence": 0.90, "library": "立项专属·历史审计台账风险库"},
    "14-用户角色权限知识库.md": {"article": "角色权限与立项边界", "effective": "现行有效", "confidence": 0.94, "library": "立项专属·用户角色权限知识库"},
    "15-审计人员负荷台账库.md": {"article": "人员负荷与资源校验", "effective": "现行有效", "confidence": 0.89, "library": "立项专属·审计人员负荷台账库"},
    "16-审计立项标准表单模板库.md": {"article": "字段 Schema 与来源规则", "effective": "现行有效", "confidence": 0.96, "library": "立项专属·标准表单模板库"},
    "17-审计岗位人员匹配知识库.md": {"article": "岗位匹配规则", "effective": "现行有效", "confidence": 0.91, "library": "立项专属·岗位人员匹配库"},
    "18-立项风险溯源知识库.md": {"article": "历史问题与行业风险索引", "effective": "现行有效", "confidence": 0.90, "library": "立项专属·风险溯源库"},
}


def _stage_for_source(source):
    for prefix, stage in _STAGE_BY_PREFIX.items():
        if source.startswith(prefix):
            return stage
    return "field"  # 未标注的文档默认归入现场实施


def _meta_for(source):
    m = _DOC_META.get(source, {})
    return {
        "article": m.get("article", "全文"),
        "effective": m.get("effective", "现行有效"),
        "confidence": m.get("confidence", 0.9),
        "library": m.get("library", "公共·审计法规库"),
    }


def _is_expired(meta):
    """过期（失效）知识不召回：生效期标注「废止/失效/过期」视为过期"""
    eff = str(meta.get("effective", ""))
    return any(k in eff for k in ("废止", "失效", "过期"))


def _split_text(text, source, max_len=500, overlap=50):
    """把文档按段落切片，超长段落按滑窗再切；每条片段携带元数据"""
    chunks = []
    paragraphs = [p.strip() for p in text.split("\n\n") if p.strip()]
    for i, para in enumerate(paragraphs):
        base = {
            "source": source,
            "id": f"{source}#{i}",
            "stage": _stage_for_source(source),
            **_meta_for(source),
        }
        if len(para) <= max_len:
            chunks.append({**base, "text": para})
        else:
            step = max_len - overlap
            for j in range(0, len(para), step):
                chunks.append({**base, "id": f"{source}#{i}-{j}", "text": para[j:j + max_len]})
    return chunks


def _template_text(path):
    """Extract text from production templates without adding a heavy office runtime.

    DOCX is OOXML and XLSX is handled by openpyxl when available. Legacy .doc/.xls
    are deliberately skipped from indexing until converted, but remain visible in
    the manifest so a missing template cannot be mistaken for a generated one.
    """
    suffix = os.path.splitext(path)[1].lower()
    if suffix == ".docx":
        try:
            with zipfile.ZipFile(path) as zf:
                xml = zf.read("word/document.xml").decode("utf-8", "ignore")
            return re.sub(r"<[^>]+>", " ", xml).replace("&amp;", "&").strip()
        except Exception:
            return ""
    if suffix == ".xlsx":
        try:
            from openpyxl import load_workbook
            wb = load_workbook(path, read_only=True, data_only=True)
            rows = []
            for ws in wb.worksheets:
                for row in ws.iter_rows(values_only=True):
                    vals = [str(v).strip() for v in row if v is not None and str(v).strip()]
                    if vals:
                        rows.append(" | ".join(vals))
            return "\n".join(rows)
        except Exception:
            return ""
    return ""


def _template_stage(path):
    rel = os.path.relpath(path, TEMPLATE_DIR).replace(os.sep, "/")
    if rel.startswith("审前准备/"):
        return "preaudit"
    if rel.startswith("现场实施/"):
        return "field"
    if rel.startswith("审计报告/"):
        return "report"
    return "field"


def load_documents():
    """加载文件知识与已发布的管理资产并切片。数据库不可用时保守回退文件目录。"""
    all_chunks = []
    for filename in sorted(os.listdir(KNOWLEDGE_DIR)):
        if not filename.endswith(".md"):
            continue
        path = os.path.join(KNOWLEDGE_DIR, filename)
        with open(path, "r", encoding="utf-8") as f:
            all_chunks.extend(_split_text(f.read(), filename))
    # 真实制式模板：解析文本进入同一索引，来源保留绝对路径和阶段。
    if os.path.isdir(TEMPLATE_DIR):
        for root, _, files in os.walk(TEMPLATE_DIR):
            for filename in sorted(files):
                if filename.startswith(".") or not filename.lower().endswith((".docx", ".xlsx")):
                    continue
                path = os.path.join(root, filename)
                text = _template_text(path)
                if text:
                    source = f"template:{os.path.relpath(path, TEMPLATE_DIR).replace(os.sep, '/') }"
                    chunks = _split_text(text, source)
                    for chunk in chunks:
                        chunk["stage"] = _template_stage(path)
                        chunk["library"] = "阶段专属·真实制式模板库"
                        chunk["template_path"] = path
                    all_chunks.extend(chunks)
    # 管理端发布的知识资产进入同一检索运行时；只加载 enabled/active 资产。
    try:
        import db
        for item in db.load_knowledge_assets():
            if item.get("status") not in ("enabled", "active"):
                continue
            source = f"managed:{item['asset_id']}:{item.get('title', '')}"
            all_chunks.extend(_split_text(item.get("content", ""), source))
    except Exception:
        pass
    return all_chunks


class BM25Retriever:
    """基于 BM25 的关键词检索器（支持阶段过滤、阈值过滤、过期过滤）"""

    def __init__(self, chunks):
        self.chunks = chunks
        tokenized = [jieba.lcut(c["text"]) for c in chunks]
        self.bm25 = BM25Okapi(tokenized)

    def retrieve(self, query, top_k=None, stage=None, threshold=None):
        """返回与 query 最相关的前 top_k 个片段。

        stage: 若指定，只返回该阶段片段。
        threshold: 归一化相似度阈值，低于此不召回（默认取 RETRIEVAL_CONFIG）。
        归一化：raw / max_raw，保证最高分片段恒为 1.0，阈值过滤只裁掉弱匹配。
        """
        k = top_k if top_k is not None else RETRIEVAL_CONFIG["top_k"]
        thr = threshold if threshold is not None else RETRIEVAL_CONFIG["threshold"]
        tokens = jieba.lcut(query)
        scores = self.bm25.get_scores(tokens)

        # 第一遍：候选集（正相关 + 未过期 + 阶段匹配）
        candidates = []
        for i in range(len(scores)):
            raw = scores[i]
            if raw <= 0:  # 过滤完全不相关的片段，避免乱引法规
                continue
            if _is_expired(self.chunks[i]):  # 过期不召回
                continue
            if stage is not None and self.chunks[i].get("stage") != stage:
                continue
            candidates.append(i)

        if not candidates:
            return []

        # 第二遍：在候选集内归一化相似度，按阈值过滤（最高分片段恒为 1.0）
        max_raw = max(float(scores[i]) for i in candidates)
        results = []
        for i in candidates:
            norm = (float(scores[i]) / max_raw) if max_raw > 0 else 0.0
            if norm < thr:
                continue
            results.append({**self.chunks[i], "score": round(float(scores[i]), 2), "similarity": round(norm, 3)})

        results.sort(key=lambda r: -r["score"])
        return results[:k]


def knowledge_manifest():
    """两级知识库结构 + 检索配置 + 示例条目清单，供 /api/knowledge 渲染"""
    public_libs = ["审计法规库", "风控规则库", "行业案例库", "提示词模板库"]
    stage_libs = ["法规库", "业务规则库", "模板库", "案例库", "工具库", "术语库"]
    docs = []
    for filename in sorted(os.listdir(KNOWLEDGE_DIR)):
        if not filename.endswith(".md"):
            continue
        meta = _meta_for(filename)
        docs.append({
            "source": filename,
            "article": meta["article"],
            "effective": meta["effective"],
            "confidence": meta["confidence"],
            "library": meta["library"],
            "level": "公共" if meta["library"].startswith("公共") else "阶段专属",
            "stage": _stage_for_source(filename),
        })
    templates = []
    if os.path.isdir(TEMPLATE_DIR):
        for root, _, files in os.walk(TEMPLATE_DIR):
            for filename in sorted(files):
                if filename.startswith("."):
                    continue
                path = os.path.join(root, filename)
                suffix = os.path.splitext(filename)[1].lower()
                templates.append({
                    "source": f"template:{os.path.relpath(path, TEMPLATE_DIR).replace(os.sep, '/')}",
                    "name": filename,
                    "path": path,
                    "stage": _template_stage(path),
                    "parseable": suffix in (".docx", ".xlsx"),
                    "status": "indexed" if suffix in (".docx", ".xlsx") else "pending_conversion",
                })
    external = {"root": EXTERNAL_KB_DIR, "enabled": os.getenv("AUDIT_EXTERNAL_KB_ENABLED", "false").lower() == "true", "files": 0}
    if os.path.isdir(EXTERNAL_KB_DIR):
        external["files"] = sum(1 for root, _, files in os.walk(EXTERNAL_KB_DIR)
                                 for name in files if not name.startswith("."))
    return {
        "public": {"name": "公共知识库", "libraries": public_libs,
                   "count": sum(1 for d in docs if d["level"] == "公共")},
        "stage": {"name": "阶段 Agent 专属知识库", "libraries": stage_libs,
                  "count": sum(1 for d in docs if d["level"] == "阶段专属")},
        "retrieval": {
            "top_k": RETRIEVAL_CONFIG["top_k"],
            "threshold": RETRIEVAL_CONFIG["threshold"],
            "requirements": ["来源", "条款位置", "生效期", "置信度", "过期不可召回", "引用须有依据"],
        },
        "docs": docs,
        "templates": templates,
        "template_dir": TEMPLATE_DIR,
        "external_knowledge": external,
    }
